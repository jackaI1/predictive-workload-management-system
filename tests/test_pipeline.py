"""In-memory sheet/workbook fixtures: no credentials, private files or API calls."""
from copy import deepcopy
from datetime import date
from io import BytesIO
import importlib
from types import SimpleNamespace
from unittest.mock import Mock
import pytest
import openpyxl
import gspread
import expander
import snapshotter
import ClickUpQA


class Sheet:
    def __init__(self, rows, number):
        self.rows = deepcopy(rows)
        self.id = number
        self.reads = 0

    def get_all_values(self):
        self.reads += 1
        return deepcopy(self.rows)

    def insert_rows(self, values, row, **kwargs):
        self.rows[row-1:row-1] = [[str(c) for c in r] for r in values]

    def update(self, range_name, values, **kwargs):
        import re
        col, row = re.match(r'([A-Z]+)(\d+)', range_name).groups()
        start = openpyxl.utils.column_index_from_string(col)-1
        target = self.rows[int(row)-1]
        while len(target) < start + len(values[0]):
            target.append('')
        target[start:start+len(values[0])] = [str(c) for c in values[0]]

    def update_cells(self, cells, **kwargs):
        for c in cells:
            self.rows[c.row-1][c.col-1] = str(c.value)


class Workbook:
    def __init__(self, rows):
        self.sheets = [Sheet(r, i) for i, r in enumerate(rows)]
        self.sheet1 = self.sheets[0]
        for sheet in self.sheets:
            sheet.spreadsheet = self

    def worksheets(self):
        return self.sheets

    def batch_update(self, body):
        for request in body['requests']:
            r = request['deleteDimension']['range']
            del self.sheets[r['sheetId']].rows[r['startIndex']:r['endIndex']]


def parent(title='D12345 AA1 TO 2', start='2026-10-19', duration='2', task_id='synthetic-parent'):
    return ['2026-10-19', title, '', '', '', '', duration, '', '', '', start, '2026-10-20', '', task_id]


def child(name):
    return ['', '', name, '1', '1000', '1.0', '', '.5', '', '2026-10-17', '2026-10-19']


def schedule():
    return [['', '', 'MODEL', '', '', 'ARCH DWG RECEIVED', 'FLOOR RECEIVED', 'TRUSS RECEIVED'],
            ['', '', '', '', ''],
            ['', '', 'AA1', 'A', '1000', '2026-10-15', '2026-10-16', '2026-10-14'],
            ['', '', 'AA2', 'A/B', '3000', '2026-10-16', '2026-10-16', '2026-10-14']]


def run_pipeline(monkeypatch, workbook):
    monkeypatch.setattr(expander, 'get_google_credentials', lambda: None)
    monkeypatch.setattr(expander.gspread, 'authorize', lambda _: SimpleNamespace(open_by_url=lambda _: workbook))
    monkeypatch.setattr(expander, 'build', lambda *a, **k: None)
    monkeypatch.setattr(expander, 'load_difficulty_tiers', lambda *a: [
        {'min':0, 'max':2500, 'base_difficulty':1.0},
        {'min':2500, 'max':float('inf'), 'base_difficulty':1.5}])
    monkeypatch.setattr(expander, 'list_spreadsheets_in_folder', lambda *a: {'D12345 synthetic': 'fixture'})
    data = schedule()
    monkeypatch.setattr(expander, 'load_and_parse_schedule', lambda *a: (data, 5, 6, 7))
    expander.main()


def test_pipeline_kpis_business_days_start_dates_and_repeat_convergence(monkeypatch, capsys):
    book = Workbook([[['header'], parent()], [['header']], [['header']]])
    run_pipeline(monkeypatch, book)
    assert 'Error' not in capsys.readouterr().out
    assert book.sheet1.rows[2][9] == '2026-10-17'
    assert book.sheet1.rows[3][5] == '1.6'
    assert sorted(r[2] for r in book.sheets[1].rows[1:]) == ['0.5', '0.5', '0.8', '0.8']
    assert {r[1] for r in book.sheets[1].rows[1:]} == {'2026-10-19', '2026-10-20'}
    assert {r[2] for r in book.sheets[2].rows[1:]} == {'0.5'}
    before = [deepcopy(s.rows) for s in book.sheets]
    run_pipeline(monkeypatch, book)
    assert 'Error' not in capsys.readouterr().out
    assert [s.rows for s in book.sheets] == before
    assert book.sheets[1].reads == book.sheets[2].reads == 2


def test_scope_removal_updates_surviving_child_without_shifting_wrong_row(monkeypatch, capsys):
    rows = [['header'], parent('D12345 AA2'), child('AA1'), child('AA2')]
    hist = [['header'], ['D12345 - AA1', '2026-10-19', '.5']]
    book = Workbook([rows, hist, hist])
    run_pipeline(monkeypatch, book)
    assert 'Error' not in capsys.readouterr().out
    assert len(book.sheet1.rows) == 3
    assert book.sheet1.rows[2][2] == 'AA2'
    assert book.sheet1.rows[2][5] == '1.6'
    assert all('AA1' not in r[0] for s in book.sheets[1:] for r in s.rows)


def test_later_dated_parent_wins_in_either_order(monkeypatch, capsys):
    old = parent('D12345 AA1', '2026-10-19', '2', 'old')
    new = parent('D12345 AA1', '2026-10-20', '4', 'new')
    for parents in ([old, new], [new, old]):
        book = Workbook([[['header'], *parents], [['header']], [['header']]])
        run_pipeline(monkeypatch, book)
        assert 'Error' not in capsys.readouterr().out
        assert {r[2] for r in book.sheets[1].rows[1:]} == {'0.25'}
        owner = None
        for row in book.sheet1.rows[1:]:
            if row[1]:
                owner = row[13]
            if row[2]:
                assert owner == 'new'
        before = [deepcopy(s.rows) for s in book.sheets]
        run_pipeline(monkeypatch, book)
        assert [s.rows for s in book.sheets] == before


def test_all_resource_identifiers_accept_environment(monkeypatch):
    for key in ('LIST_ID', 'PROGRESSION_STATUS_FIELD_ID', 'GOOGLE_SHEET_URL'):
        monkeypatch.setenv(key, 'synthetic-config')
    for module in (snapshotter, ClickUpQA):
        importlib.reload(module)
        try:
            for key in ('LIST_ID', 'PROGRESSION_STATUS_FIELD_ID', 'GOOGLE_SHEET_URL'):
                assert getattr(module, key) == 'synthetic-config'
        finally:
            for key in ('LIST_ID', 'PROGRESSION_STATUS_FIELD_ID', 'GOOGLE_SHEET_URL'):
                monkeypatch.delenv(key)
            importlib.reload(module)
            for key in ('LIST_ID', 'PROGRESSION_STATUS_FIELD_ID', 'GOOGLE_SHEET_URL'):
                monkeypatch.setenv(key, 'synthetic-config')


@pytest.mark.parametrize('module', [expander, snapshotter, ClickUpQA])
def test_sheet_backoff_and_non_retryable_errors(monkeypatch, module):
    response = Mock()
    response.json.return_value = {'error': {'code':429, 'message':'synthetic quota', 'status':'RESOURCE_EXHAUSTED'}}
    quota = gspread.exceptions.APIError(response)
    monkeypatch.setattr(module.time, 'sleep', Mock())
    monkeypatch.setattr(module.random, 'uniform', lambda *a: 0)
    operation = Mock(side_effect=[quota, quota, 'ok'])
    assert module.execute_with_retry(operation) == 'ok'
    sleeps = [c.args[0] for c in module.time.sleep.call_args_list]
    assert len(sleeps) == 2 and sleeps[1] == sleeps[0]*2
    operation = Mock(side_effect=ValueError('invalid'))
    with pytest.raises(ValueError):
        module.execute_with_retry(operation)
    assert operation.call_count == 1


def test_tracker_backoff_pagination_and_incomplete_stream_abort(monkeypatch):
    quota = SimpleNamespace(status_code=429)
    ok = SimpleNamespace(status_code=200, json=lambda: {'tasks':[{'id':'synthetic'}]})
    done = SimpleNamespace(status_code=200, json=lambda: {'tasks':[]})
    session = Mock()
    session.get.side_effect = [quota, ok, done]
    monkeypatch.setattr(snapshotter.requests, 'Session', lambda: session)
    monkeypatch.setattr(snapshotter.time, 'sleep', Mock())
    assert snapshotter.get_clickup_tasks() == [{'id':'synthetic', '_page':0}]
    assert session.get.call_count == 3
    session.get.side_effect = [SimpleNamespace(status_code=403, text='synthetic denial')]
    with pytest.raises(RuntimeError, match='Sync aborted'):
        snapshotter.get_clickup_tasks()
    session.request.side_effect = [quota, done]
    assert ClickUpQA.clickup_request(session, 'GET', 'https://example.invalid') is done


def test_workbook_download_and_merged_cells_are_parsed():
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in schedule():
        ws.append(row)
    ws.merge_cells('F3:F4')
    data = BytesIO()
    wb.save(data)
    service = Mock()
    service.files.return_value.get_media.return_value.execute.return_value = data.getvalue()
    rows, arch, floor, truss = expander.load_and_parse_schedule('synthetic', service)
    assert (arch, floor, truss) == (5, 6, 7)
    assert rows[2][5] == rows[3][5] == '2026-10-15'


def test_remote_tiers_fallback_and_elevation_scoring():
    client = Mock()
    client.open_by_url.return_value.sheet1.get_all_values.return_value = [
        ['sqft', 'difficulty'], ['2500', '1.7'], ['0', '1.0']]
    tiers = expander.load_difficulty_tiers(client, 'synthetic')
    assert expander.parse_sq_ft_and_difficulty('2,800/3000', expander.parse_elevation_count('A/B/C'), tiers) == ('3000', 1.9)
    client.open_by_url.side_effect = RuntimeError('offline fixture')
    fallback = expander.load_difficulty_tiers(client, 'synthetic')
    assert expander.parse_sq_ft_and_difficulty('4000', 2, fallback) == ('4000', 2.1)


def test_empty_remote_tiers_use_documented_fallback():
    client = Mock()
    client.open_by_url.return_value.sheet1.get_all_values.return_value = [['sqft', 'difficulty']]
    tiers = expander.load_difficulty_tiers(client, 'synthetic')
    assert expander.parse_sq_ft_and_difficulty('4000', 1, tiers) == ('4000', 2.0)


def test_block_start_date_waits_for_missing_member_and_rolls_up_latest_date():
    from expander import derive_start_dates
    def item(name, block=False, floor=date(2026,10,16)):
        return dict(model_output=name, is_block=block, arch_dt=date(2026,10,15),
                    floor_dt=floor, truss_dt=date(2026,10,14), final_start_date='')
    items = [item('AA1'), item('AA2', floor=None), item('BLOCK 1', True)]
    derive_start_dates(items)
    assert items[-1]['final_start_date'] == ''
    items[1]['floor_dt'] = date(2026,10,20)
    derive_start_dates(items)
    assert items[-1]['final_start_date'] == '2026-10-21'


def test_equal_date_ownership_is_deterministic_and_master_header_is_preserved(monkeypatch, capsys):
    a = parent('D12345 AA1', task_id='synthetic-a', duration='2')
    b = parent('D12345 AA1', task_id='synthetic-b', duration='4')
    for parents in ([a, b], [b, a]):
        header = ['Snapshot', 'Task', 'Model']
        book = Workbook([[header, *parents], [['header']], [['header']]])
        run_pipeline(monkeypatch, book)
        assert 'Error' not in capsys.readouterr().out
        assert book.sheet1.rows[0] == header
        assert {r[2] for r in book.sheets[1].rows[1:]} == {'0.25'}
